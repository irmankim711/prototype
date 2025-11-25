import os
import sys
from flask import Flask
from docx import Document

# Add the application directory to the python path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), 'backend')))

# Mock the app context
app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.root_path = os.path.abspath(os.path.join(os.path.dirname(__file__), 'backend/app'))

# Create a dummy template
template_dir = os.path.join(os.path.dirname(__file__), 'backend/templates')
os.makedirs(template_dir, exist_ok=True)
template_path = os.path.join(template_dir, 'test_template.docx')

doc = Document()
doc.add_paragraph('Hello {{ name }}!')
doc.save(template_path)

print(f"Created test template at {template_path}")

# Mock the models import to avoid database dependency
import sys
from unittest.mock import MagicMock
sys.modules['app.models'] = MagicMock()
sys.modules['app.models.TemplateModel'] = MagicMock()

# Import the service
from app.services.export_service import export_service

# Run the export
with app.app_context():
    try:
        print("Running export...")
        result = export_service.export(
            template_id='test_template',
            data_source={'name': 'World'},
            formats=['docx']
        )
        
        docx_filename = result.filenames['docx']
        print(f"Export successful. Filename: {docx_filename}")
        
        # Check the content
        # The service puts it in backend/uploads/reports because root_path is backend/app
        reports_dir = os.path.join(os.path.dirname(__file__), 'backend/uploads/reports')
        output_path = os.path.join(reports_dir, docx_filename)
        
        print(f"Checking output at: {output_path}")
        
        if os.path.exists(output_path):
            doc = Document(output_path)
            text = [p.text for p in doc.paragraphs]
            print(f"Output content: {text}")
            
            if 'Hello World!' in text:
                print("SUCCESS: Template was rendered.")
            elif 'Hello {{ name }}!' in text:
                print("FAILURE: Template was NOT rendered (raw template copied).")
            else:
                print(f"UNKNOWN: Unexpected content: {text}")
        else:
            print(f"Error: Output file not found at {output_path}")
            # List directory to see what's there
            if os.path.exists(reports_dir):
                print(f"Contents of {reports_dir}: {os.listdir(reports_dir)}")
            else:
                print(f"Directory {reports_dir} does not exist")
            
    except Exception as e:
        print(f"Error during export: {e}")
        import traceback
        traceback.print_exc()
