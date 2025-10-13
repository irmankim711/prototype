"""
Convert docxtemplater template to Jinja2 format for docxtpl
Converts {variable} to {{variable}}
"""
from docx import Document
import re
import sys

def convert_template(input_path, output_path):
    """Convert template tags from {var} to {{var}}"""
    doc = Document(input_path)
    conversions = 0

    # Convert in paragraphs
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            original_text = run.text
            # Replace {tag} with {{tag}} but avoid already converted tags
            # First, temporarily mark already-converted tags
            text = original_text.replace('{{', '\x00DOUBLE_OPEN\x00')
            text = text.replace('}}', '\x00DOUBLE_CLOSE\x00')

            # Now convert single braces
            text = re.sub(r'\{([^}]+)\}', r'{{\1}}', text)

            # Restore any already-converted tags
            text = text.replace('\x00DOUBLE_OPEN\x00', '{{')
            text = text.replace('\x00DOUBLE_CLOSE\x00', '}}')

            if text != original_text:
                run.text = text
                conversions += 1
                print(f"Converted: {original_text[:50]} -> {text[:50]}")

    # Convert in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        original_text = run.text
                        # Same conversion logic
                        text = original_text.replace('{{', '\x00DOUBLE_OPEN\x00')
                        text = text.replace('}}', '\x00DOUBLE_CLOSE\x00')
                        text = re.sub(r'\{([^}]+)\}', r'{{\1}}', text)
                        text = text.replace('\x00DOUBLE_OPEN\x00', '{{')
                        text = text.replace('\x00DOUBLE_CLOSE\x00', '}}')

                        if text != original_text:
                            run.text = text
                            conversions += 1

    # Save converted template
    doc.save(output_path)
    print(f"\n✅ Conversion complete!")
    print(f"   Total conversions: {conversions}")
    print(f"   Output: {output_path}")

    return conversions

if __name__ == "__main__":
    input_file = "backend/templates/report_templates/04- LAPORAN FU _ PUNCAK ALAM (1).docx"
    output_file = "backend/templates/04- LAPORAN FU _ PUNCAK ALAM_converted.docx"

    print(f"Converting template...")
    print(f"Input:  {input_file}")
    print(f"Output: {output_file}\n")

    try:
        conversions = convert_template(input_file, output_file)

        # Verify conversion
        print("\n🔍 Verifying conversion...")
        doc = Document(output_file)
        tags = set()

        for p in doc.paragraphs:
            tags.update(re.findall(r'\{\{([^}]+)\}\}', p.text))

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        tags.update(re.findall(r'\{\{([^}]+)\}\}', p.text))

        print(f"\nJinja2 tags found in converted template:")
        for tag in sorted(tags):
            print(f"  {{{{ {tag} }}}}")
        print(f"\nTotal unique Jinja2 tags: {len(tags)}")

    except Exception as e:
        print(f"❌ Error: {e}")
        sys.exit(1)
