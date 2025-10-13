"""
Manually find and fix the specific cells with line breaks
"""
from docx import Document
import re

def find_and_fix_cells(input_path, output_path):
    """Find cells with multiline tags and fix them"""
    doc = Document(input_path)
    fixes = 0

    print("Searching for cells with multiline tags...\n")

    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                cell_text = cell.text

                # Check if this cell contains the problematic tags
                if '{{' in cell_text and '\n' in cell_text and '}}' in cell_text:
                    # Check for our specific problematic patterns
                    if 'KEHADIRAH' in cell_text or 'KEHADIRAN' in cell_text:
                        print(f"Found at Table {table_idx}, Row {row_idx}, Cell {cell_idx}")
                        print(f"  Original: {repr(cell_text)}")

                        # Clear the cell and rewrite with clean text
                        for paragraph in cell.paragraphs:
                            paragraph.clear()

                        # Add clean text - rename the variables to remove line breaks
                        if 'KEHADIRAH' in cell_text:
                            clean_text = "{{ KEHADIRAH_SABTU }}"
                        elif 'KEHADIRAN' in cell_text and 'AHAD' in cell_text:
                            clean_text = "{{ KEHADIRAN_AHAD }}"
                        else:
                            # Fallback: try to clean it generically
                            clean_text = re.sub(r'\s+', ' ', cell_text.replace('\n', ' ').replace('\r', ' '))

                        cell.paragraphs[0].add_run(clean_text)
                        print(f"  Fixed:    {repr(clean_text)}\n")
                        fixes += 1

    doc.save(output_path)
    print(f"✅ Fixed {fixes} cells with multiline tags")
    return fixes

if __name__ == "__main__":
    input_file = "backend/templates/04- LAPORAN FU _ PUNCAK ALAM_converted.docx"
    output_file = "backend/templates/04- LAPORAN FU _ PUNCAK ALAM_final.docx"

    fixes = find_and_fix_cells(input_file, output_file)

    # Verify
    print("\n🔍 Verification...")
    doc = Document(output_file)
    tags = set()

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for tag in re.findall(r'\{\{\s*([^}]+?)\s*\}\}', cell.text):
                    tags.add(tag.strip())

    for p in doc.paragraphs:
        for tag in re.findall(r'\{\{\s*([^}]+?)\s*\}\}', p.text):
            tags.add(tag.strip())

    print(f"\n✅ All Jinja2 tags in final template:")
    for tag in sorted(tags):
        # Check for line breaks
        if '\n' in tag or '\r' in tag:
            print(f"  ⚠️  {{{{ {repr(tag)} }}}}")
        else:
            print(f"  {{{{ {tag} }}}}")

    print(f"\nTotal: {len(tags)} unique tags")

    clean_tags = [t for t in tags if '\n' not in t and '\r' not in t]
    if len(clean_tags) == len(tags):
        print("\n✅ SUCCESS! All tags are clean with no line breaks!")
    else:
        print(f"\n⚠️  {len(tags) - len(clean_tags)} tags still have line breaks")
