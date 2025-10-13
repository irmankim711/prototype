"""
Fix multiline tags in template
Removes line breaks from within Jinja2 tags
"""
from docx import Document
import re

def fix_multiline_tags(input_path, output_path):
    """Fix tags that have line breaks inside them"""
    doc = Document(input_path)
    fixes = 0

    # Fix in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        original_text = run.text

                        # Find and fix multiline tags
                        # Pattern: {{ TAG\nOTHER_TEXT\n }}
                        text = re.sub(
                            r'\{\{([^}]*)\n([^}]*)\}\}',
                            lambda m: '{{' + m.group(1).strip() + ' ' + m.group(2).strip() + '}}',
                            original_text
                        )

                        # Also handle tags with line breaks at the boundaries
                        text = re.sub(r'\{\{\s*\n\s*', '{{ ', text)
                        text = re.sub(r'\s*\n\s*\}\}', ' }}', text)

                        if text != original_text:
                            run.text = text
                            fixes += 1
                            print(f"Fixed: {repr(original_text)} -> {repr(text)}")

    # Fix in paragraphs
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            original_text = run.text
            text = re.sub(
                r'\{\{([^}]*)\n([^}]*)\}\}',
                lambda m: '{{' + m.group(1).strip() + ' ' + m.group(2).strip() + '}}',
                original_text
            )
            text = re.sub(r'\{\{\s*\n\s*', '{{ ', text)
            text = re.sub(r'\s*\n\s*\}\}', ' }}', text)

            if text != original_text:
                run.text = text
                fixes += 1

    doc.save(output_path)
    print(f"\n✅ Fixed {fixes} multiline tags")
    print(f"   Output: {output_path}")

    return fixes

if __name__ == "__main__":
    input_file = "backend/templates/04- LAPORAN FU _ PUNCAK ALAM_converted.docx"
    output_file = "backend/templates/04- LAPORAN FU _ PUNCAK ALAM_final.docx"

    print("Fixing multiline tags...\n")
    fixes = fix_multiline_tags(input_file, output_file)

    # Verify
    print("\n🔍 Verifying all tags...")
    doc = Document(output_file)
    tags = set()

    for p in doc.paragraphs:
        for tag in re.findall(r'\{\{([^}]+)\}\}', p.text):
            tags.add(tag.strip())

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for tag in re.findall(r'\{\{([^}]+)\}\}', cell.text):
                    tags.add(tag.strip())

    print(f"\nAll Jinja2 tags in final template:")
    for tag in sorted(tags):
        print(f"  {{{{ {tag} }}}}")
    print(f"\nTotal unique tags: {len(tags)}")
