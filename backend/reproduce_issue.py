import sys
import os
from pathlib import Path
import openpyxl
from jinja2 import Template
import traceback

# Add the parent directory to Python path
sys.path.insert(0, str(Path(__file__).parent))

from app.services.form_automation import FormAutomationService
from app import create_app

def create_test_excel_file():
    """Create a test Excel file with sample data"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sample Data"
    
    # Headers
    headers = ["Region", "Quarter", "Revenue", "Profit Margin", "Customer Count"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # Sample data
    data = [
        ["North", "Q1", 120000, 0.15, 450],
        ["North", "Q2", 135000, 0.18, 520],
    ]
    
    for row, row_data in enumerate(data, 2):
        for col, value in enumerate(row_data, 1):
            ws.cell(row=row, column=col, value=value)
    
    # Save to test file
    test_file = Path(__file__).parent / "reproduce_data.xlsx"
    wb.save(test_file)
    return test_file

def create_test_template():
    """Create a test report template"""
    template_content = """
# {{ report_title }}
{% for region in regions %}
### {{ region.name }}
- Revenue: ${{ region.revenue }}
{% endfor %}
"""
    
    template_file = Path(__file__).parent / "reproduce_template.jinja"
    with open(template_file, 'w', encoding='utf-8') as f:
        f.write(template_content)
    
    return template_file

def reproduce():
    print("🚀 Starting Reproduction Script")
    
    app = create_app('testing')
    
    with app.app_context():
        try:
            excel_file = create_test_excel_file()
            template_file = create_test_template()
            
            form_automation = FormAutomationService()
            
            print("\n📝 Testing report generation...")
            # We call the method that failed
            report_result = form_automation.generate_report_from_excel(
                excel_path=str(excel_file),
                template_path=str(template_file)
            )
            
            if report_result['success']:
                print("✅ Report generation successful")
            else:
                print(f"❌ Report generation failed: {report_result.get('error')}")
                # If error info is returned, print it
                if 'error_type' in report_result:
                    print(f"Error Type: {report_result['error_type']}")
            
        except Exception as e:
            print(f"❌ Exception caught: {str(e)}")
            traceback.print_exc()
        finally:
            if 'excel_file' in locals() and excel_file.exists():
                excel_file.unlink()
            if 'template_file' in locals() and template_file.exists():
                template_file.unlink()

if __name__ == "__main__":
    reproduce()
