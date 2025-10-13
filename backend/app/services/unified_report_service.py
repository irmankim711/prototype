"""
Unified Report Generation Service
Consolidated service for generating PDF and DOCX reports with data from various sources
"""

import os
import json
import uuid
from datetime import datetime
from typing import Dict, List, Any, Optional, Tuple
from pathlib import Path
import logging

# Report generation libraries
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .. import db
from ..models import Report, ReportTemplate, User
from ..core.exceptions import ReportGenerationError

logger = logging.getLogger(__name__)

class UnifiedReportService:
    """Unified service for generating reports with consistent API"""

    def __init__(self, upload_folder: str = None):
        """Initialize the service with upload directory"""
        self.upload_folder = upload_folder or os.path.join(os.getcwd(), 'uploads', 'reports')
        self.ensure_upload_directory()
        self.styles = getSampleStyleSheet()
        self.custom_styles = self._create_custom_styles()

    def ensure_upload_directory(self):
        """Ensure the upload directory exists"""
        Path(self.upload_folder).mkdir(parents=True, exist_ok=True)

    def _create_custom_styles(self):
        """Create custom paragraph styles for reports"""
        custom_styles = {}

        # Title style
        custom_styles['Title'] = ParagraphStyle(
            'CustomTitle',
            parent=self.styles['Title'],
            fontSize=20,
            spaceAfter=30,
            textColor=colors.darkblue,
            alignment=TA_CENTER
        )

        # Heading style
        custom_styles['Heading'] = ParagraphStyle(
            'CustomHeading',
            parent=self.styles['Heading1'],
            fontSize=14,
            spaceAfter=12,
            textColor=colors.darkblue,
            spaceBefore=20
        )

        # Body style
        custom_styles['Body'] = ParagraphStyle(
            'CustomBody',
            parent=self.styles['Normal'],
            fontSize=11,
            spaceAfter=12
        )

        return custom_styles

    def generate_timestamp_filename(self, prefix: str, extension: str) -> str:
        """Generate a unique filename with timestamp"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        unique_id = str(uuid.uuid4())[:8]
        return f"{prefix}_{timestamp}_{unique_id}.{extension}"

    def create_report(self, title: str, description: str, template_id: int = None,
                     data_source: Dict = None, user_id: int = None) -> Report:
        """
        Create a new report record in the database

        Args:
            title: Report title
            description: Report description
            template_id: Template ID to use
            data_source: Source data for the report
            user_id: User creating the report

        Returns:
            Report: The created report object
        """
        try:
            report = Report(
                title=title,
                description=description,
                template_id=template_id,
                data_source=data_source,
                generation_status='pending',
                created_by=str(user_id) if user_id else None,
                report_type='standard'
            )

            db.session.add(report)
            db.session.commit()

            logger.info(f"Created report {report.id}: {title}")
            return report

        except Exception as e:
            db.session.rollback()
            logger.error(f"Failed to create report: {e}")
            raise ReportGenerationError(f"Failed to create report: {str(e)}")

    def generate_pdf_report(self, report: Report, data: Dict[str, Any]) -> str:
        """
        Generate a PDF report

        Args:
            report: Report object
            data: Data to include in the report

        Returns:
            str: Path to the generated PDF file
        """
        try:
            # Update report status
            report.update_status('generating', processing_notes='Starting PDF generation')
            db.session.commit()

            # Generate filename
            filename = self.generate_timestamp_filename(f"report_{report.id}", "pdf")
            file_path = os.path.join(self.upload_folder, filename)

            # Create PDF document
            doc = SimpleDocTemplate(file_path, pagesize=A4)
            story = []

            # Add title
            title = Paragraph(report.title or "Generated Report", self.custom_styles['Title'])
            story.append(title)
            story.append(Spacer(1, 20))

            # Add generation info
            generation_info = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            story.append(Paragraph(generation_info, self.custom_styles['Body']))
            story.append(Spacer(1, 20))

            # Add description if available
            if report.description:
                story.append(Paragraph("Description:", self.custom_styles['Heading']))
                story.append(Paragraph(report.description, self.custom_styles['Body']))
                story.append(Spacer(1, 20))

            # Add data content
            if data:
                story.append(Paragraph("Report Data:", self.custom_styles['Heading']))

                # Handle different data types
                if isinstance(data, dict):
                    for key, value in data.items():
                        if isinstance(value, (str, int, float)):
                            story.append(Paragraph(f"<b>{key}:</b> {value}", self.custom_styles['Body']))
                        elif isinstance(value, list) and value:
                            story.append(Paragraph(f"<b>{key}:</b>", self.custom_styles['Body']))
                            # Create table for list data
                            if all(isinstance(item, dict) for item in value):
                                table_data = self._prepare_table_data(value)
                                if table_data:
                                    table = Table(table_data)
                                    table.setStyle(TableStyle([
                                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                                        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                                        ('FONTSIZE', (0, 0), (-1, 0), 12),
                                        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                                        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                                        ('GRID', (0, 0), (-1, -1), 1, colors.black)
                                    ]))
                                    story.append(table)
                            else:
                                for item in value[:10]:  # Limit to first 10 items
                                    story.append(Paragraph(f"• {item}", self.custom_styles['Body']))

                story.append(Spacer(1, 20))

            # Build PDF
            doc.build(story)

            # Update report with file info
            file_size = os.path.getsize(file_path)
            report.file_path = file_path
            report.file_size = file_size
            report.file_format = 'pdf'
            report.download_url = f"/api/reports/download/{filename}"
            report.update_status('completed', processing_notes='PDF generation completed successfully')

            db.session.commit()

            logger.info(f"Generated PDF report {report.id}: {file_path}")
            return file_path

        except Exception as e:
            report.update_status('failed', error_message=str(e))
            db.session.commit()
            logger.error(f"Failed to generate PDF report {report.id}: {e}")
            raise ReportGenerationError(f"PDF generation failed: {str(e)}")

    def generate_docx_report(self, report: Report, data: Dict[str, Any]) -> str:
        """
        Generate a DOCX report

        Args:
            report: Report object
            data: Data to include in the report

        Returns:
            str: Path to the generated DOCX file
        """
        try:
            # Update report status
            report.update_status('generating', processing_notes='Starting DOCX generation')
            db.session.commit()

            # Generate filename
            filename = self.generate_timestamp_filename(f"report_{report.id}", "docx")
            file_path = os.path.join(self.upload_folder, filename)

            # Create DOCX document
            doc = Document()

            # Add title
            title = doc.add_heading(report.title or "Generated Report", 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Add generation info
            doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph("")

            # Add description if available
            if report.description:
                doc.add_heading("Description", level=1)
                doc.add_paragraph(report.description)
                doc.add_paragraph("")

            # Add data content
            if data:
                doc.add_heading("Report Data", level=1)

                # Handle different data types
                if isinstance(data, dict):
                    for key, value in data.items():
                        if isinstance(value, (str, int, float)):
                            p = doc.add_paragraph()
                            p.add_run(f"{key}: ").bold = True
                            p.add_run(str(value))
                        elif isinstance(value, list) and value:
                            p = doc.add_paragraph()
                            p.add_run(f"{key}:").bold = True
                            # Create table for list data
                            if all(isinstance(item, dict) for item in value):
                                table_data = self._prepare_table_data(value)
                                if table_data and len(table_data) > 1:
                                    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                                    table.style = 'Table Grid'

                                    for i, row in enumerate(table_data):
                                        for j, cell_value in enumerate(row):
                                            cell = table.cell(i, j)
                                            cell.text = str(cell_value)
                                            if i == 0:  # Header row
                                                cell.paragraphs[0].runs[0].bold = True
                            else:
                                for item in value[:10]:  # Limit to first 10 items
                                    doc.add_paragraph(f"• {item}", style='List Bullet')

                doc.add_paragraph("")

            # Save document
            doc.save(file_path)

            # Update report with file info
            file_size = os.path.getsize(file_path)
            report.file_path = file_path
            report.file_size = file_size
            report.file_format = 'docx'
            report.download_url = f"/api/reports/download/{filename}"
            report.update_status('completed', processing_notes='DOCX generation completed successfully')

            db.session.commit()

            logger.info(f"Generated DOCX report {report.id}: {file_path}")
            return file_path

        except Exception as e:
            report.update_status('failed', error_message=str(e))
            db.session.commit()
            logger.error(f"Failed to generate DOCX report {report.id}: {e}")
            raise ReportGenerationError(f"DOCX generation failed: {str(e)}")

    def _prepare_table_data(self, data_list: List[Dict]) -> List[List[str]]:
        """Prepare data for table creation"""
        if not data_list:
            return []

        # Get headers from first item
        headers = list(data_list[0].keys())
        table_data = [headers]

        # Add data rows
        for item in data_list[:20]:  # Limit to first 20 rows
            row = [str(item.get(header, '')) for header in headers]
            table_data.append(row)

        return table_data

    def get_templates(self) -> List[Dict]:
        """Get all available report templates"""
        try:
            templates = ReportTemplate.query.filter_by(is_active=True).all()
            return [
                {
                    'id': t.id,
                    'name': t.name,
                    'description': t.description,
                    'template_type': t.template_type,
                    'parameters': t.parameters
                }
                for t in templates
            ]
        except Exception as e:
            logger.error(f"Failed to get templates: {e}")
            return []

    def get_report(self, report_id: int) -> Optional[Report]:
        """Get a report by ID"""
        try:
            return Report.query.get(report_id)
        except Exception as e:
            logger.error(f"Failed to get report {report_id}: {e}")
            return None

    def delete_report(self, report_id: int) -> bool:
        """Delete a report and its associated files"""
        try:
            report = Report.query.get(report_id)
            if not report:
                return False

            # Delete file if it exists
            if report.file_path and os.path.exists(report.file_path):
                os.remove(report.file_path)
                logger.info(f"Deleted file: {report.file_path}")

            # Delete from database
            db.session.delete(report)
            db.session.commit()

            logger.info(f"Deleted report {report_id}")
            return True

        except Exception as e:
            db.session.rollback()
            logger.error(f"Failed to delete report {report_id}: {e}")
            return False


# Create a singleton instance
unified_report_service = UnifiedReportService()