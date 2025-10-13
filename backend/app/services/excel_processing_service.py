"""
Excel Processing Service for Template-Based Report Generation
Handles Excel file upload, parsing, and data extraction
"""
import os
import logging
import hashlib
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime
from pathlib import Path
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename
import pandas as pd
import openpyxl
from openpyxl import load_workbook
from flask import current_app

from ..models.template_models import ExcelUpload
from .. import db

logger = logging.getLogger(__name__)

class ExcelProcessingService:
    """Service for handling Excel file operations"""
    
    def __init__(self):
        self.upload_folder = Path(current_app.root_path).parent / 'uploads' / 'excel'
        self.upload_folder.mkdir(parents=True, exist_ok=True)
        
        # Allowed file extensions
        self.allowed_extensions = {'.xlsx', '.xls', '.xlsm'}
        
        # Maximum file size (50MB)
        self.max_file_size = 50 * 1024 * 1024
        
        logger.info("Excel processing service initialized")
    
    def upload_file(self, file: FileStorage, user_id: int) -> Dict[str, Any]:
        """
        Upload and process Excel file
        
        Args:
            file: FileStorage object from Flask request
            user_id: User ID uploading the file
            
        Returns:
            Upload result dictionary
        """
        try:
            # Validate file
            validation_result = self._validate_file(file)
            if not validation_result['valid']:
                return {
                    'success': False,
                    'error': 'File validation failed',
                    'details': validation_result['errors']
                }
            
            # Generate secure filename
            original_filename = file.filename
            filename = self._generate_secure_filename(original_filename, user_id)
            file_path = self.upload_folder / filename
            
            # Save file
            file.save(str(file_path))
            
            # Get file info
            file_size = file_path.stat().st_size
            mime_type = self._get_mime_type(file_path)
            
            # Parse Excel file
            parse_result = self.parse_excel(str(file_path))
            
            # Create database record
            excel_upload = ExcelUpload(
                filename=filename,
                original_filename=original_filename,
                file_path=str(file_path),
                file_size=file_size,
                mime_type=mime_type,
                uploaded_by=user_id,
                sheet_names=parse_result.get('sheet_names', []),
                column_headers=parse_result.get('column_headers', {}),
                row_count=parse_result.get('total_rows', 0),
                processed=parse_result.get('success', False),
                processing_errors=parse_result.get('errors', [])
            )
            
            db.session.add(excel_upload)
            db.session.commit()
            
            logger.info(f"Excel file uploaded successfully: {filename} by user {user_id}")
            
            return {
                'success': True,
                'upload_id': excel_upload.id,
                'filename': filename,
                'original_filename': original_filename,
                'file_size': file_size,
                'parse_result': parse_result
            }
            
        except Exception as e:
            db.session.rollback()
            logger.error(f"Error uploading Excel file: {e}")
            
            # Clean up file if it was saved
            if 'file_path' in locals() and file_path.exists():
                try:
                    file_path.unlink()
                except:
                    pass
            
            return {
                'success': False,
                'error': 'File upload failed',
                'message': str(e)
            }
    
    def parse_excel(self, file_path: str) -> Dict[str, Any]:
        """
        Parse Excel file and extract metadata
        
        Args:
            file_path: Path to Excel file
            
        Returns:
            Parse result dictionary
        """
        try:
            if not os.path.exists(file_path):
                return {
                    'success': False,
                    'error': 'File not found',
                    'file_path': file_path
                }
            
            # Load workbook
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            
            result = {
                'success': True,
                'file_path': file_path,
                'sheet_names': workbook.sheetnames,
                'column_headers': {},
                'sheet_info': {},
                'total_rows': 0,
                'errors': [],
                'warnings': []
            }
            
            # Process each sheet
            for sheet_name in workbook.sheetnames:
                try:
                    sheet = workbook[sheet_name]
                    
                    # Get sheet dimensions
                    max_row = sheet.max_row
                    max_col = sheet.max_column
                    
                    # Extract column headers (first row)
                    headers = []
                    if max_row > 0:
                        for col in range(1, max_col + 1):
                            cell_value = sheet.cell(row=1, column=col).value
                            headers.append(str(cell_value) if cell_value is not None else f'Column_{col}')
                    
                    # Store sheet information
                    result['column_headers'][sheet_name] = headers
                    result['sheet_info'][sheet_name] = {
                        'max_row': max_row,
                        'max_col': max_col,
                        'data_rows': max_row - 1 if max_row > 0 else 0,  # Excluding header row
                        'column_count': len(headers)
                    }
                    
                    result['total_rows'] += max_row - 1 if max_row > 0 else 0
                    
                    logger.debug(f"Processed sheet '{sheet_name}': {max_row} rows, {max_col} columns")
                    
                except Exception as sheet_error:
                    error_msg = f"Error processing sheet '{sheet_name}': {str(sheet_error)}"
                    result['errors'].append(error_msg)
                    logger.warning(error_msg)
            
            workbook.close()
            
            # Add warnings for common issues
            if result['total_rows'] == 0:
                result['warnings'].append('No data rows found in any sheet')
            
            if len(result['sheet_names']) > 5:
                result['warnings'].append(f'File contains {len(result["sheet_names"])} sheets - consider using fewer sheets for better performance')
            
            logger.info(f"Excel file parsed successfully: {len(result['sheet_names'])} sheets, {result['total_rows']} total rows")
            
            return result
            
        except Exception as e:
            logger.error(f"Error parsing Excel file {file_path}: {e}")
            return {
                'success': False,
                'error': 'Excel parsing failed',
                'message': str(e),
                'file_path': file_path
            }
    
    def get_sheets(self, file_path: str) -> List[str]:
        """
        Get list of sheet names from Excel file
        
        Args:
            file_path: Path to Excel file
            
        Returns:
            List of sheet names
        """
        try:
            workbook = load_workbook(file_path, read_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
            return sheet_names
            
        except Exception as e:
            logger.error(f"Error getting sheets from {file_path}: {e}")
            return []
    
    def extract_data(self, file_path: str, sheet: str = None, max_rows: int = None) -> pd.DataFrame:
        """
        Extract data from Excel file as pandas DataFrame
        
        Args:
            file_path: Path to Excel file
            sheet: Sheet name (uses first sheet if None)
            max_rows: Maximum number of rows to read
            
        Returns:
            pandas DataFrame with the data
        """
        try:
            # Read Excel file
            if sheet:
                df = pd.read_excel(file_path, sheet_name=sheet, nrows=max_rows)
            else:
                df = pd.read_excel(file_path, nrows=max_rows)
            
            # Clean column names
            df.columns = df.columns.astype(str)
            df.columns = [col.strip() for col in df.columns]
            
            # Remove completely empty rows
            df = df.dropna(how='all')
            
            # Convert data types
            df = self._optimize_data_types(df)
            
            logger.info(f"Extracted data: {len(df)} rows, {len(df.columns)} columns from {file_path}")
            
            return df
            
        except Exception as e:
            logger.error(f"Error extracting data from {file_path}: {e}")
            return pd.DataFrame()  # Return empty DataFrame on error
    
    def infer_data_types(self, data: pd.DataFrame) -> Dict[str, str]:
        """
        Infer data types for DataFrame columns
        
        Args:
            data: pandas DataFrame
            
        Returns:
            Dictionary mapping column names to inferred types
        """
        type_mapping = {}
        
        for column in data.columns:
            series = data[column].dropna()
            
            if len(series) == 0:
                type_mapping[column] = 'string'
                continue
            
            # Check for numeric types
            if pd.api.types.is_numeric_dtype(series):
                if pd.api.types.is_integer_dtype(series):
                    type_mapping[column] = 'integer'
                else:
                    type_mapping[column] = 'float'
            
            # Check for datetime
            elif pd.api.types.is_datetime64_any_dtype(series):
                type_mapping[column] = 'datetime'
            
            # Check for boolean
            elif pd.api.types.is_bool_dtype(series):
                type_mapping[column] = 'boolean'
            
            # Try to infer datetime from strings
            elif series.dtype == 'object':
                # Sample a few values to check for date patterns
                sample_values = series.head(10).astype(str)
                date_like_count = 0
                
                for value in sample_values:
                    if self._looks_like_date(value):
                        date_like_count += 1
                
                if date_like_count >= len(sample_values) * 0.7:  # 70% look like dates
                    type_mapping[column] = 'date'
                else:
                    type_mapping[column] = 'string'
            
            else:
                type_mapping[column] = 'string'
        
        return type_mapping
    
    def validate_data_structure(self, data: pd.DataFrame, template_vars: List[Dict[str, Any]]) -> Dict[str, Any]:
        """
        Validate data structure against template requirements
        
        Args:
            data: pandas DataFrame with Excel data
            template_vars: List of template variable definitions
            
        Returns:
            Validation result dictionary
        """
        validation_result = {
            'valid': True,
            'errors': [],
            'warnings': [],
            'suggestions': [],
            'column_mapping': {},
            'missing_variables': [],
            'extra_columns': []
        }
        
        try:
            # Get required template variables
            required_vars = [var for var in template_vars if var.get('required', False)]
            optional_vars = [var for var in template_vars if not var.get('required', False)]
            
            data_columns = list(data.columns)
            template_var_names = [var.get('name', '') for var in template_vars]
            
            # Check for missing required variables
            for var in required_vars:
                var_name = var.get('name', '')
                if not self._find_matching_column(var_name, data_columns):
                    validation_result['missing_variables'].append(var_name)
                    validation_result['errors'].append(f"Required variable '{var_name}' not found in Excel data")
                    validation_result['valid'] = False
            
            # Suggest column mappings
            for var in template_vars:
                var_name = var.get('name', '')
                matching_column = self._find_matching_column(var_name, data_columns)
                if matching_column:
                    validation_result['column_mapping'][var_name] = matching_column
            
            # Identify extra columns
            mapped_columns = list(validation_result['column_mapping'].values())
            validation_result['extra_columns'] = [col for col in data_columns if col not in mapped_columns]
            
            # Add suggestions
            if validation_result['extra_columns']:
                validation_result['suggestions'].append(f"Found {len(validation_result['extra_columns'])} unmapped columns that could be used for additional template variables")
            
            if len(data) == 0:
                validation_result['warnings'].append("Excel file contains no data rows")
            elif len(data) > 10000:
                validation_result['warnings'].append(f"Large dataset ({len(data)} rows) may impact performance")
            
            logger.info(f"Data structure validation completed - Valid: {validation_result['valid']}")
            
        except Exception as e:
            logger.error(f"Error validating data structure: {e}")
            validation_result['valid'] = False
            validation_result['errors'].append(f"Validation error: {str(e)}")
        
        return validation_result
    
    def get_data_preview(self, file_path: str, sheet: str = None, rows: int = 10) -> Dict[str, Any]:
        """
        Get preview of Excel data
        
        Args:
            file_path: Path to Excel file
            sheet: Sheet name
            rows: Number of rows to preview
            
        Returns:
            Preview data dictionary
        """
        try:
            df = self.extract_data(file_path, sheet, max_rows=rows)
            
            if df.empty:
                return {
                    'success': False,
                    'error': 'No data found'
                }
            
            # Convert DataFrame to dictionary format
            preview_data = {
                'success': True,
                'columns': list(df.columns),
                'data': df.to_dict('records'),
                'total_columns': len(df.columns),
                'preview_rows': len(df),
                'data_types': self.infer_data_types(df)
            }
            
            return preview_data
            
        except Exception as e:
            logger.error(f"Error getting data preview: {e}")
            return {
                'success': False,
                'error': str(e)
            }
    
    def _validate_file(self, file: FileStorage) -> Dict[str, Any]:
        """Validate uploaded file"""
        result = {'valid': True, 'errors': []}
        
        # Check if file is provided
        if not file or not file.filename:
            result['valid'] = False
            result['errors'].append('No file provided')
            return result
        
        # Check file extension
        file_ext = Path(file.filename).suffix.lower()
        if file_ext not in self.allowed_extensions:
            result['valid'] = False
            result['errors'].append(f'Invalid file type. Allowed: {", ".join(self.allowed_extensions)}')
        
        # Check file size (approximate)
        if hasattr(file, 'content_length') and file.content_length:
            if file.content_length > self.max_file_size:
                result['valid'] = False
                result['errors'].append(f'File too large. Maximum size: {self.max_file_size // (1024*1024)}MB')
        
        return result
    
    def _generate_secure_filename(self, original_filename: str, user_id: int) -> str:
        """Generate secure filename with timestamp and user ID"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        name_part = secure_filename(Path(original_filename).stem)
        extension = Path(original_filename).suffix.lower()
        
        # Create hash for uniqueness
        hash_input = f"{user_id}_{original_filename}_{timestamp}"
        file_hash = hashlib.md5(hash_input.encode()).hexdigest()[:8]
        
        return f"{name_part}_{timestamp}_{user_id}_{file_hash}{extension}"
    
    def _get_mime_type(self, file_path: Path) -> str:
        """Get MIME type based on file extension"""
        extension = file_path.suffix.lower()
        mime_types = {
            '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            '.xls': 'application/vnd.ms-excel',
            '.xlsm': 'application/vnd.ms-excel.sheet.macroEnabled.12'
        }
        return mime_types.get(extension, 'application/octet-stream')
    
    def _optimize_data_types(self, df: pd.DataFrame) -> pd.DataFrame:
        """Optimize DataFrame data types for memory efficiency"""
        for column in df.columns:
            if df[column].dtype == 'object':
                # Try to convert to numeric
                numeric_series = pd.to_numeric(df[column], errors='coerce')
                if not numeric_series.isna().all():
                    df[column] = numeric_series
                
                # Try to convert to datetime
                elif df[column].astype(str).apply(self._looks_like_date).any():
                    try:
                        df[column] = pd.to_datetime(df[column], errors='coerce')
                    except:
                        pass
        
        return df
    
    def _looks_like_date(self, value: str) -> bool:
        """Check if string value looks like a date"""
        if not isinstance(value, str) or len(value.strip()) < 6:
            return False
        
        # Common date patterns
        date_patterns = [
            r'\d{4}-\d{1,2}-\d{1,2}',  # YYYY-MM-DD
            r'\d{1,2}/\d{1,2}/\d{4}',  # MM/DD/YYYY or DD/MM/YYYY
            r'\d{1,2}-\d{1,2}-\d{4}',  # MM-DD-YYYY or DD-MM-YYYY
            r'\d{4}/\d{1,2}/\d{1,2}',  # YYYY/MM/DD
        ]
        
        import re
        for pattern in date_patterns:
            if re.match(pattern, value.strip()):
                return True
        
        return False
    
    def _find_matching_column(self, var_name: str, columns: List[str]) -> Optional[str]:
        """Find matching column for template variable"""
        var_name_lower = var_name.lower()
        
        # Exact match (case insensitive)
        for col in columns:
            if col.lower() == var_name_lower:
                return col
        
        # Partial match
        for col in columns:
            if var_name_lower in col.lower() or col.lower() in var_name_lower:
                return col
        
        # Fuzzy match (remove common separators)
        var_clean = var_name_lower.replace('_', '').replace('-', '').replace(' ', '')
        for col in columns:
            col_clean = col.lower().replace('_', '').replace('-', '').replace(' ', '')
            if var_clean == col_clean:
                return col
        
        return None

# Global service instance
excel_processing_service = ExcelProcessingService()