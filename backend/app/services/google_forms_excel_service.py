"""
Google Forms Excel Export Service
Specialized service for exporting Google Forms responses to Excel format
"""

import os
import json
import uuid
import time
from datetime import datetime, timedelta
from typing import Dict, List, Any, Optional, Tuple
from pathlib import Path
import logging

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.drawing.image import Image
import pandas as pd

from ..services.google_forms_service import google_forms_service
from ..core.exceptions import ReportGenerationError as ExportError

logger = logging.getLogger(__name__)

class GoogleFormsExcelService:
    """Service for exporting Google Forms responses to Excel with advanced formatting"""

    def __init__(self, upload_folder: str = None):
        self.upload_folder = upload_folder or os.path.join(os.getcwd(), 'static', 'exports')
        self.ensure_upload_directory()

    def ensure_upload_directory(self):
        """Ensure the upload directory exists"""
        Path(self.upload_folder).mkdir(parents=True, exist_ok=True)

    def generate_timestamp_filename(self, prefix: str, form_title: str = None) -> str:
        """Generate a unique filename with timestamp"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        unique_id = str(uuid.uuid4())[:8]

        if form_title:
            clean_title = "".join(c for c in form_title if c.isalnum() or c in (' ', '-', '_')).rstrip()
            clean_title = clean_title.replace(' ', '_')[:30]
            return f"{prefix}_{clean_title}_{timestamp}_{unique_id}.xlsx"
        else:
            return f"{prefix}_google_form_{timestamp}_{unique_id}.xlsx"

    def export_google_form_to_excel(self, user_id: str, form_id: str, options: Dict[str, Any] = None) -> Dict[str, Any]:
        """Export Google Form responses to Excel"""
        start_time = time.time()

        try:
            if not google_forms_service or not google_forms_service.is_enabled():
                raise ExportError("Google Forms service is not available")

            options = options or {}
            include_analytics = options.get('include_analytics', True)
            excel_options = options.get('excel_options', {})
            max_responses = options.get('max_responses', 1000)
            use_ai_enhancement = options.get('use_ai_enhancement', False)

            logger.info(f"Starting Google Forms export for form {form_id} by user {user_id} (AI Enhanced: {use_ai_enhancement})")

            # Get form responses with analysis
            form_data = google_forms_service.get_form_responses_for_automated_report(
                user_id, form_id
            )

            if not form_data.get('success', False):
                raise ExportError(f"Failed to fetch Google Form data: {form_data.get('error', 'Unknown error')}")

            responses = form_data.get('responses', [])
            form_info = form_data.get('form_info', {})
            analysis = form_data.get('analysis', {})

            # Limit responses if specified
            if len(responses) > max_responses:
                responses = responses[:max_responses]
                logger.warning(f"Limited export to {max_responses} responses out of {len(form_data.get('responses', []))}")

            # Use AI-enhanced export if requested
            ai_fallback_reason = None
            if use_ai_enhancement:
                try:
                    return self._export_with_ai_enhancement(
                        user_id, form_id, responses, form_info, analysis,
                        excel_options, include_analytics, start_time
                    )
                except ExportError as e:
                    # Log the AI enhancement failure and fall back to standard export
                    ai_fallback_reason = str(e)
                    logger.warning(f"AI enhancement failed, falling back to standard export: {ai_fallback_reason}")
                    # Continue with standard export below
                except Exception as e:
                    # Catch any other errors and fall back
                    ai_fallback_reason = f"Unexpected error: {str(e)}"
                    logger.error(f"Unexpected AI enhancement error: {ai_fallback_reason}", exc_info=True)
                    # Continue with standard export below

            # Standard export (existing code)
            # Generate filename and create workbook
            filename = self.generate_timestamp_filename("google_forms_export", form_info.get('title'))
            file_path = os.path.join(self.upload_folder, filename)

            wb = openpyxl.Workbook()

            # Remove default sheet and create our sheets
            wb.remove(wb.active)

            # Create main responses sheet
            responses_sheet = wb.create_sheet("Responses")
            self._export_responses_to_sheet(responses_sheet, responses, form_info)

            # Create analytics sheet if requested
            if include_analytics and analysis:
                analytics_sheet = wb.create_sheet("Analytics")
                self._export_analytics_to_sheet(analytics_sheet, analysis, form_info)

            # Create form info sheet
            if excel_options.get('include_form_schema', True):
                info_sheet = wb.create_sheet("Form Info")
                self._export_form_info_to_sheet(info_sheet, form_info, form_data)

            # Apply professional formatting if requested
            if excel_options.get('formatting', 'professional') == 'professional':
                self._apply_professional_formatting(wb)

            # Save workbook
            wb.save(file_path)

            # Calculate metrics
            file_size = os.path.getsize(file_path)
            generation_time = time.time() - start_time

            # Generate download URL
            download_url = f"/api/google-forms/forms/{form_id}/download-excel/{filename}"

            logger.info(f"Google Forms Excel export completed: {file_path} ({file_size} bytes)")

            result = {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'download_url': download_url,
                'file_size': file_size,
                'responses_count': len(responses),
                'generation_time': generation_time,
                'form_info': form_info,
                'data_quality_score': self._calculate_data_quality_score(responses),
                'ai_enhanced': False
            }

            # Add warning if AI enhancement was requested but failed
            if ai_fallback_reason:
                result['warning'] = f"AI enhancement failed: {ai_fallback_reason}. Standard export generated instead."
                result['ai_requested_but_failed'] = True
                logger.warning(f"Returning standard export with AI fallback warning: {ai_fallback_reason}")

            return result

        except Exception as e:
            logger.error(f"Error exporting Google Form {form_id} to Excel: {str(e)}")
            return {
                'success': False,
                'error': f"Failed to export Google Form to Excel: {str(e)}"
            }

    def _export_responses_to_sheet(self, ws, responses: List[Dict], form_info: Dict):
        """Export Google Forms responses to Excel worksheet"""
        # CRITICAL: Add debugging to identify empty responses
        logger.info(f"_export_responses_to_sheet called with {len(responses)} responses")

        if not responses:
            ws.cell(row=1, column=1, value="No responses found")
            logger.warning("No responses to export - writing 'No responses found' message")
            return

        # Extract all unique question titles from responses
        all_questions = set()
        for idx, response in enumerate(responses):
            response_answers = response.get('answers', {})
            if not response_answers:
                logger.warning(f"Response {idx} (ID: {response.get('response_id', 'unknown')}) has empty answers dict!")
            all_questions.update(response_answers.keys())

        question_list = sorted(list(all_questions))

        # CRITICAL: Check if we found any questions
        if not question_list:
            logger.error("❌ CRITICAL BUG: No questions found in any response!")
            logger.error("This means all responses have empty 'answers' dictionaries")
            logger.error(f"Sample response structure: {json.dumps(responses[0] if responses else {}, indent=2)}")

            # Still write headers with metadata only
            ws.cell(row=1, column=1, value="Response ID")
            ws.cell(row=1, column=2, value="Submission Time")
            ws.cell(row=1, column=3, value="ERROR: No question data found")

            # Write response IDs to show we received data
            for row_idx, response in enumerate(responses, 2):
                ws.cell(row=row_idx, column=1, value=response.get('response_id', ''))
                ws.cell(row=row_idx, column=2, value=response.get('create_time', ''))
                ws.cell(row=row_idx, column=3, value="Empty answers dict")
            return

        logger.info(f"Found {len(question_list)} unique questions: {question_list[:5]}...")

        # Prepare headers
        headers = ['Response ID', 'Submission Time', 'Last Modified']
        headers.extend(question_list)

        # Write headers with styling
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        # Write data rows
        rows_written = 0
        for row_idx, response in enumerate(responses, 2):
            col = 1

            # Response metadata
            response_id = response.get('response_id', '')
            ws.cell(row=row_idx, column=col, value=response_id)
            col += 1

            create_time = response.get('create_time', '')
            if create_time:
                try:
                    # Convert ISO format to readable datetime
                    dt = datetime.fromisoformat(create_time.replace('Z', '+00:00'))
                    ws.cell(row=row_idx, column=col, value=dt.strftime('%Y-%m-%d %H:%M:%S'))
                except:
                    ws.cell(row=row_idx, column=col, value=create_time)
            col += 1

            last_submitted = response.get('last_submitted_time', '')
            if last_submitted:
                try:
                    dt = datetime.fromisoformat(last_submitted.replace('Z', '+00:00'))
                    ws.cell(row=row_idx, column=col, value=dt.strftime('%Y-%m-%d %H:%M:%S'))
                except:
                    ws.cell(row=row_idx, column=col, value=last_submitted)
            col += 1

            # Response answers - CRITICAL: Iterate through questions in same order as headers
            answers = response.get('answers', {})
            answers_written = 0

            for question in question_list:
                answer = answers.get(question, '')

                # Handle different answer types
                if isinstance(answer, list):
                    # List of choices - join with commas
                    answer = ', '.join([str(item) for item in answer if item])
                elif isinstance(answer, dict):
                    # Complex structure - convert to JSON
                    answer = json.dumps(answer, ensure_ascii=False)
                elif answer is None or answer == '':
                    answer = ''
                else:
                    answer = str(answer)

                ws.cell(row=row_idx, column=col, value=answer)
                if answer:  # Count non-empty answers
                    answers_written += 1
                col += 1

            rows_written += 1
            if row_idx == 2:  # Log first data row for debugging
                logger.info(f"First data row written: Response ID={response_id}, Answers={answers_written}/{len(question_list)}")

        logger.info(f"✅ Successfully wrote {rows_written} data rows with {len(question_list)} question columns")

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _export_analytics_to_sheet(self, ws, analysis: Dict, form_info: Dict):
        """Export analytics data to Excel worksheet"""
        row = 1

        # Title
        title_cell = ws.cell(row=row, column=1, value="Google Forms Analytics Report")
        title_cell.font = Font(bold=True, size=16, color="2F5597")
        row += 2

        # Form Information
        ws.cell(row=row, column=1, value="Form Information").font = Font(bold=True, size=12)
        row += 1

        form_data = [
            ("Form Title", form_info.get('title', 'N/A')),
            ("Form ID", form_info.get('id', 'N/A')),
            ("Total Questions", form_info.get('total_questions', 0)),
            ("Analysis Date", datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        ]

        for label, value in form_data:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=str(value))
            row += 1

        row += 1

        # Response Statistics
        completion_stats = analysis.get('completion_stats', {})
        if completion_stats:
            ws.cell(row=row, column=1, value="Response Statistics").font = Font(bold=True, size=12)
            row += 1

            stats_data = [
                ("Total Responses", completion_stats.get('total_responses', 0)),
                ("Completion Rate", f"{completion_stats.get('completion_rate', 0)}%"),
                ("First Response", completion_stats.get('first_response', 'N/A')),
                ("Last Response", completion_stats.get('last_response', 'N/A'))
            ]

            for label, value in stats_data:
                ws.cell(row=row, column=1, value=label).font = Font(bold=True)
                ws.cell(row=row, column=2, value=str(value))
                row += 1

            row += 1

        # Question Insights
        field_analysis = analysis.get('field_analysis', {})
        if field_analysis:
            ws.cell(row=row, column=1, value="Question Analysis").font = Font(bold=True, size=12)
            row += 1

            # Headers for question analysis
            headers = ["Question", "Type", "Response Count", "Average Score", "Most Common Answer"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
            row += 1

            for question, data in field_analysis.items():
                ws.cell(row=row, column=1, value=question)
                ws.cell(row=row, column=2, value=data.get('type', 'N/A'))

                if data['type'] == 'rating' and 'statistics' in data:
                    stats = data['statistics']
                    ws.cell(row=row, column=3, value=stats.get('count', 0))
                    ws.cell(row=row, column=4, value=round(stats.get('mean', 0), 2))
                elif data['type'] == 'feedback':
                    ws.cell(row=row, column=3, value=len(data.get('responses', [])))
                    ws.cell(row=row, column=4, value="N/A")

                # Most common answer would require additional analysis
                ws.cell(row=row, column=5, value="N/A")
                row += 1

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _export_form_info_to_sheet(self, ws, form_info: Dict, form_data: Dict):
        """Export form information to Excel worksheet"""
        row = 1

        # Title
        title_cell = ws.cell(row=row, column=1, value="Google Form Information")
        title_cell.font = Font(bold=True, size=16, color="2F5597")
        row += 2

        # Basic Information
        basic_info = [
            ("Form Title", form_info.get('title', 'N/A')),
            ("Form Description", form_info.get('description', 'N/A')),
            ("Form ID", form_info.get('id', 'N/A')),
            ("Published URL", form_info.get('published_url', 'N/A')),
            ("Total Questions", form_info.get('total_questions', 0)),
            ("Form Type", form_info.get('form_type', 'Google Form')),
            ("Export Date", datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ("Data Source", form_data.get('data_source', 'google_forms_api')),
            ("Retrieved At", form_data.get('retrieved_at', 'N/A'))
        ]

        for label, value in basic_info:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=str(value))
            row += 1

        row += 1

        # Questions Information
        questions = form_data.get('questions', {})
        if questions:
            ws.cell(row=row, column=1, value="Questions Details").font = Font(bold=True, size=12)
            row += 1

            # Headers
            headers = ["Question ID", "Question Title", "Question Type", "Required"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
            row += 1

            for question_id, question_data in questions.items():
                ws.cell(row=row, column=1, value=question_id)
                ws.cell(row=row, column=2, value=question_data.get('title', 'N/A'))
                ws.cell(row=row, column=3, value=question_data.get('type', 'N/A'))
                ws.cell(row=row, column=4, value="Yes" if question_data.get('required', False) else "No")
                row += 1

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 60)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _apply_professional_formatting(self, wb):
        """Apply professional formatting to the workbook"""
        # Define color scheme
        primary_color = "4472C4"
        secondary_color = "E7E6E6"
        accent_color = "70AD47"

        for sheet in wb.worksheets:
            # Apply borders to all cells with data
            max_row = sheet.max_row
            max_col = sheet.max_column

            if max_row > 0 and max_col > 0:
                thin_border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )

                for row in range(1, max_row + 1):
                    for col in range(1, max_col + 1):
                        sheet.cell(row=row, column=col).border = thin_border

                # Freeze first row
                sheet.freeze_panes = "A2"

    def _calculate_data_quality_score(self, responses: List[Dict]) -> float:
        """Calculate a data quality score for the responses"""
        if not responses:
            return 0.0

        total_questions = 0
        answered_questions = 0

        for response in responses:
            answers = response.get('answers', {})
            response_questions = len(answers)
            response_answered = len([a for a in answers.values() if a and str(a).strip()])

            total_questions += response_questions
            answered_questions += response_answered

        if total_questions == 0:
            return 0.0

        quality_score = answered_questions / total_questions
        return min(quality_score, 1.0)

    def _export_with_ai_enhancement(
        self,
        user_id: str,
        form_id: str,
        responses: List[Dict],
        form_info: Dict,
        analysis: Dict,
        excel_options: Dict,
        include_analytics: bool,
        start_time: float
    ) -> Dict[str, Any]:
        """
        Export Google Form with AI-enhanced Excel formatting

        Args:
            user_id: User identifier (for logging and tracking)
            form_id: Google Form ID being exported
            responses: List of form responses
            form_info: Form metadata
            analysis: Pre-computed analysis from Google Forms service (available for future use)
            excel_options: Export configuration options
            include_analytics: Whether to include analytics sheet
            start_time: Export start timestamp

        Returns:
            Dict containing export result with file paths and metadata
        """
        try:
            # Import AI Excel service
            from app.services.ai_enhanced_excel_service import ai_excel_service

            logger.info(f"User {user_id} initiating AI-enhanced Excel generation for form {form_id}")

            # Note: The 'analysis' parameter contains pre-computed insights from Google Forms service
            # It's kept for potential future enhancements where we can pass existing analysis to AI
            # Currently, the AI service performs its own analysis via Claude API

            # Convert Google Forms responses to flat data structure for AI service
            data = []
            for response in responses:
                flat_response = {
                    'Response ID': response.get('response_id', ''),
                    'Submission Time': response.get('create_time', ''),
                    'Last Modified': response.get('last_submitted_time', ''),
                }

                # Add all answer fields
                answers = response.get('answers', {})
                for question, answer in answers.items():
                    # Convert lists/dicts to strings
                    if isinstance(answer, list):
                        flat_response[question] = ', '.join([str(item) for item in answer if item])
                    elif isinstance(answer, dict):
                        flat_response[question] = json.dumps(answer, ensure_ascii=False)
                    else:
                        flat_response[question] = str(answer) if answer else ''

                data.append(flat_response)

            # Prepare title with form info
            title = form_info.get('title', 'Google Form Responses')

            # Prepare AI enhancement options
            ai_options = {
                'include_summary': include_analytics,
                'include_charts': excel_options.get('include_charts', True),
                'include_pivot': excel_options.get('include_pivot', False),
            }

            # Generate AI-enhanced Excel
            file_path, file_size = ai_excel_service.generate_enhanced_excel(
                data=data,
                title=title,
                options=ai_options
            )

            # Calculate metrics
            generation_time = time.time() - start_time

            # Extract filename from path
            filename = os.path.basename(file_path)

            # Move file to export directory if needed
            target_path = os.path.join(self.upload_folder, filename)
            if file_path != target_path:
                import shutil
                shutil.move(file_path, target_path)
                file_path = target_path

            # Generate download URL
            download_url = f"/api/google-forms/forms/{form_id}/download-excel/{filename}"

            logger.info(f"AI-enhanced Google Forms Excel export completed: {file_path} ({file_size} bytes)")

            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'download_url': download_url,
                'file_size': file_size,
                'responses_count': len(responses),
                'generation_time': generation_time,
                'form_info': form_info,
                'data_quality_score': self._calculate_data_quality_score(responses),
                'ai_enhanced': True
            }

        except ImportError as e:
            logger.exception(f"AI Excel service not available for user {user_id}, form {form_id}: {str(e)}")
            raise ExportError("AI enhancement service not available")
        except Exception as e:
            logger.exception(f"Error in AI-enhanced export for user {user_id}, form {form_id}")
            # Re-raise to allow caller to handle fallback
            raise ExportError(f"AI-enhanced export failed: {str(e)}")

# Global instance
google_forms_excel_service = GoogleFormsExcelService()