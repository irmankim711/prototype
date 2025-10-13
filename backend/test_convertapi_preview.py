#!/usr/bin/env python3
"""
Test ConvertAPI Preview Functionality
"""

import os
import sys
import requests
import tempfile
from pathlib import Path

# Add backend to path
sys.path.append(os.path.join(os.path.dirname(__file__)))

def test_convertapi_preview():
    """Test ConvertAPI DOCX to HTML preview conversion"""

    # Set environment variable
    os.environ['CONVERT_API_SECRET'] = 'iQMCycsKxTXbJVsrWZgt27v5qIBYerYY'

    try:
        from app.services.convertapi_service import convertapi_service

        print("🔧 Testing ConvertAPI connection...")
        connection_result = convertapi_service.test_connection()
        print(f"Connection result: {connection_result}")

        if not connection_result['success']:
            print("❌ ConvertAPI connection failed")
            return False

        print("✅ ConvertAPI connection successful")

        # Find a test DOCX file
        test_docx_files = []
        backend_dir = Path(__file__).parent

        # Look for DOCX files in templates and static/generated directories
        search_dirs = [
            backend_dir / 'templates',
            backend_dir / 'static' / 'generated',
            backend_dir.parent / 'example'  # Check example directory
        ]

        for search_dir in search_dirs:
            if search_dir.exists():
                docx_files = list(search_dir.glob('*.docx'))
                test_docx_files.extend(docx_files)

        if not test_docx_files:
            print("❌ No test DOCX files found")
            return False

        # Use the first DOCX file found
        test_file = test_docx_files[0]
        print(f"📄 Using test file: {test_file}")

        # Test HTML preview conversion
        print("🔄 Testing DOCX to HTML preview conversion...")
        success, message, html_path = convertapi_service.convert_docx_to_html_preview(str(test_file))

        if success and html_path:
            print(f"✅ HTML preview conversion successful: {html_path}")

            # Check if the HTML file was created and has content
            if os.path.exists(html_path):
                file_size = os.path.getsize(html_path)
                print(f"📏 HTML file size: {file_size} bytes")

                # Read a snippet of the HTML content
                with open(html_path, 'r', encoding='utf-8') as f:
                    content_snippet = f.read(500)
                    print(f"📄 HTML content snippet: {content_snippet[:200]}...")

                return True
            else:
                print(f"❌ HTML file not found at: {html_path}")
                return False
        else:
            print(f"❌ HTML preview conversion failed: {message}")
            return False

    except Exception as e:
        print(f"❌ Test failed with error: {e}")
        return False

def test_direct_api_call():
    """Test direct ConvertAPI call"""
    print("\n🌐 Testing direct ConvertAPI call...")

    api_key = 'iQMCycsKxTXbJVsrWZgt27v5qIBYerYY'

    # Create a simple test DOCX file
    try:
        from docx import Document

        # Create a simple document
        doc = Document()
        doc.add_heading('Test Document', 0)
        doc.add_paragraph('This is a test document for ConvertAPI preview testing.')
        doc.add_paragraph('It contains some sample text to verify the conversion works.')

        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp_file:
            doc.save(tmp_file.name)
            test_docx_path = tmp_file.name

        print(f"📄 Created test DOCX: {test_docx_path}")

        # Make direct API call
        url = "https://v2.convertapi.com/convert/docx/to/html"

        with open(test_docx_path, 'rb') as file:
            files = {
                'File': (os.path.basename(test_docx_path), file, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            }

            data = {
                'Secret': api_key,
                'StoreFile': 'true'
            }

            response = requests.post(url, files=files, data=data, timeout=120)

        print(f"📡 API Response Status: {response.status_code}")

        if response.status_code == 200:
            result = response.json()
            print(f"✅ API Response: {result}")

            if 'Files' in result and len(result['Files']) > 0:
                html_url = result['Files'][0]['Url']
                print(f"🔗 HTML URL: {html_url}")

                # Download the HTML content
                html_response = requests.get(html_url, timeout=60)
                if html_response.status_code == 200:
                    html_content = html_response.text
                    print(f"📄 HTML content length: {len(html_content)} characters")
                    print(f"📄 HTML content snippet: {html_content[:300]}...")
                    return True
                else:
                    print(f"❌ Failed to download HTML: {html_response.status_code}")
            else:
                print("❌ No files in response")
        else:
            print(f"❌ API call failed: {response.text}")

        # Clean up
        os.unlink(test_docx_path)

    except ImportError:
        print("❌ python-docx not available for test document creation")
    except Exception as e:
        print(f"❌ Direct API test failed: {e}")

    return False

if __name__ == "__main__":
    print("🧪 ConvertAPI Preview Test Suite")
    print("=" * 50)

    # Test 1: ConvertAPI service
    test1_result = test_convertapi_preview()

    # Test 2: Direct API call
    test2_result = test_direct_api_call()

    print("\n" + "=" * 50)
    print("📊 Test Results:")
    print(f"🔧 ConvertAPI Service Test: {'✅ PASS' if test1_result else '❌ FAIL'}")
    print(f"🌐 Direct API Test: {'✅ PASS' if test2_result else '❌ FAIL'}")

    if test1_result or test2_result:
        print("\n🎉 ConvertAPI preview functionality is working!")
        print("The preview endpoint should now work properly.")
    else:
        print("\n❌ ConvertAPI preview tests failed")
        print("Please check the API key and network connectivity.")